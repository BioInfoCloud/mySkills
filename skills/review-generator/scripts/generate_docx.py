#!/usr/bin/env python3
"""
Generate formatted Word (.docx) documents from markdown review files.

Supports both Chinese (楷体 or 宋体) and English (Times New Roman) formatting.
Produces review_CN.docx and review_EN.docx with proper academic styling.

Usage:
    python generate_docx.py --input review_CN.md --output review_CN.docx --lang zh
    python generate_docx.py --input review_EN.md --output review_EN.docx --lang en
    python generate_docx.py --input-dir ./RNH1_review/  # auto-detect both files

Requirements:
    pip install python-docx
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml, OxmlElement
except ImportError:
    print("python-docx is required. Install with: pip install python-docx")
    sys.exit(1)


# ---------------------------------------------------------------------------
# Font configuration by language
# ---------------------------------------------------------------------------
FONT_CONFIG = {
    "zh": {
        "body_font": "楷体",
        "heading_font": "黑体",
        "body_font_ascii": "Times New Roman",
        "title_size": Pt(16),
        "h1_size": Pt(14),
        "h2_size": Pt(12),
        "h3_size": Pt(12),
        "body_size": Pt(12),
        "ref_size": Pt(10),
        "abstract_size": Pt(10),
    },
    "en": {
        "body_font": "Times New Roman",
        "heading_font": "Times New Roman",
        "body_font_ascii": "Times New Roman",
        "title_size": Pt(16),
        "h1_size": Pt(14),
        "h2_size": Pt(12),
        "h3_size": Pt(12),
        "body_size": Pt(12),
        "ref_size": Pt(10),
        "abstract_size": Pt(10),
    },
}


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------
def parse_markdown(md_text):
    """Parse markdown into a list of (type, content) tuples."""
    lines = md_text.split("\n")
    blocks = []
    i = 0
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            if in_table and table_rows:
                blocks.append(("table", table_rows))
                table_rows = []
                in_table = False
            i += 1
            continue

        # Table detection
        if stripped.startswith("|") and stripped.endswith("|"):
            # Check if next line is separator (|---|---|)
            if i + 1 < len(lines) and re.match(r"^\|[\s\-:]+\|", lines[i + 1].strip()):
                in_table = True
                # Parse header row
                header = [c.strip() for c in stripped.split("|")[1:-1]]
                table_rows.append(header)
                i += 2  # skip header and separator
                continue
            elif in_table:
                row = [c.strip() for c in stripped.split("|")[1:-1]]
                table_rows.append(row)
                i += 1
                continue

        if in_table and table_rows:
            blocks.append(("table", table_rows))
            table_rows = []
            in_table = False

        # Title (# )
        if stripped.startswith("# ") and not stripped.startswith("## "):
            blocks.append(("title", stripped[2:].strip()))
            i += 1
            continue

        # H1
        if stripped.startswith("## ") and not stripped.startswith("### "):
            blocks.append(("h1", stripped[3:].strip()))
            i += 1
            continue

        # H2
        if stripped.startswith("### ") and not stripped.startswith("#### "):
            blocks.append(("h2", stripped[4:].strip()))
            i += 1
            continue

        # H3
        if stripped.startswith("#### "):
            blocks.append(("h3", stripped[5:].strip()))
            i += 1
            continue

        # Horizontal rule (skip)
        if stripped in ("---", "***", "___"):
            i += 1
            continue

        # Regular paragraph (collect multi-line paragraphs)
        para_lines = [line]
        j = i + 1
        while j < len(lines):
            next_line = lines[j].strip()
            if (not next_line or
                next_line.startswith("#") or
                next_line.startswith("|") or
                next_line in ("---", "***", "___")):
                break
            para_lines.append(lines[j])
            j += 1
        para_text = " ".join(l.strip() for l in para_lines)
        blocks.append(("paragraph", para_text))
        i = j

    # Flush any remaining table
    if in_table and table_rows:
        blocks.append(("table", table_rows))

    return blocks


# ---------------------------------------------------------------------------
# Font helper: set both Western and East Asian fonts explicitly
# ---------------------------------------------------------------------------
def _set_run_font(run, western_font, eastasia_font, size):
    """Set font on a run, specifying both Western (ascii/hAnsi) and East Asian fonts."""
    run.font.size = size
    run.font.name = western_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), western_font)
    rfonts.set(qn('w:hAnsi'), western_font)
    rfonts.set(qn('w:eastAsia'), eastasia_font)
    rfonts.set(qn('w:cs'), western_font)


def _run_font_config(config, is_heading=False):
    """Return (western_font, eastasia_font) tuple based on language config."""
    if config["body_font"] == "宋体" or config["body_font"] == "楷体":
        western = config["body_font_ascii"]  # Times New Roman
        eastasia = config["heading_font"] if is_heading else config["body_font"]
    else:
        western = config["body_font"]  # Times New Roman
        eastasia = config["body_font"]  # Times New Roman
    return western, eastasia


# ---------------------------------------------------------------------------
# Inline formatting (bold, italic, citations)
# ---------------------------------------------------------------------------
def add_runs_with_formatting(paragraph, text, lang_config, is_reference=False):
    """Add text to paragraph, parsing **bold**, *italic*, and [N] citations."""
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|\[[\d\-,\s]+\])'
    parts = re.split(pattern, text)

    font_size = lang_config["ref_size"] if is_reference else lang_config["body_size"]
    western_font, eastasia_font = _run_font_config(lang_config)

    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif re.match(r'^\[[\d\-,\s]+\]$', part):
            run = paragraph.add_run(part)
            run.font.superscript = True
        else:
            run = paragraph.add_run(part)

        _set_run_font(run, western_font, eastasia_font, font_size)


# ---------------------------------------------------------------------------
# Table rendering
# ---------------------------------------------------------------------------
def add_table_to_doc(doc, table_data, lang_config):
    """Add a formatted table to the document."""
    if not table_data or len(table_data) < 1:
        return

    num_cols = len(table_data[0])
    table = doc.add_table(rows=len(table_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Apply table style
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else parse_xml(
        f'<w:tblPr {nsdecls("w")}></w:tblPr>'
    )

    # Add borders
    borders_xml = f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
    </w:tblBorders>'''
    tbl_pr.append(parse_xml(borders_xml))

    for row_idx, row_data in enumerate(table_data):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= num_cols:
                break
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs_with_formatting(para, cell_text, lang_config)

            # Header row bold
            if row_idx == 0:
                for run in para.runs:
                    run.bold = True
                # Shading for header
                shading = parse_xml(
                    f'<w:shd {nsdecls("w")} w:fill="2F5496" w:val="clear"/>'
                )
                cell._tc.tcPr.append(shading)
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Spacing after table
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Document generation
# ---------------------------------------------------------------------------
def generate_docx(blocks, output_path, lang="zh"):
    """Generate a formatted .docx file from parsed blocks."""
    config = FONT_CONFIG[lang]
    doc = Document()

    # Set default page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # Set default style — explicitly define both Western and East Asian fonts
    western, eastasia = _run_font_config(config)
    style = doc.styles["Normal"]
    style.font.size = config["body_size"]
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), western)
    rfonts.set(qn('w:hAnsi'), western)
    rfonts.set(qn('w:eastAsia'), eastasia)
    rfonts.set(qn('w:cs'), western)

    # Track whether we're in the references section
    in_references = False

    for block_type, content in blocks:
        if block_type == "title":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(content)
            run.bold = True
            western, eastasia = _run_font_config(config, is_heading=True)
            _set_run_font(run, western, eastasia, config["title_size"])
            para.paragraph_format.space_after = Pt(18)
            para.paragraph_format.space_before = Pt(0)

        elif block_type == "h1":
            # Check if entering references section
            lower_content = content.lower()
            if "参考文献" in content or "references" in lower_content:
                in_references = True

            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = para.add_run(content)
            run.bold = True
            western, eastasia = _run_font_config(config, is_heading=True)
            _set_run_font(run, western, eastasia, config["h1_size"])
            para.paragraph_format.space_before = Pt(18)
            para.paragraph_format.space_after = Pt(12)

        elif block_type == "h2":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = para.add_run(content)
            run.bold = True
            western, eastasia = _run_font_config(config, is_heading=True)
            _set_run_font(run, western, eastasia, config["h2_size"])
            para.paragraph_format.space_before = Pt(12)
            para.paragraph_format.space_after = Pt(6)

        elif block_type == "h3":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = para.add_run(content)
            run.bold = True
            run.italic = False
            western, eastasia = _run_font_config(config, is_heading=True)
            _set_run_font(run, western, eastasia, config["h3_size"])
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)

        elif block_type == "table":
            add_table_to_doc(doc, content, config)

        elif block_type == "paragraph":
            # Skip image-only lines; strip inline images
            if re.match(r'^!\[.*\]\(.*\)$', content.strip()):
                continue
            content = re.sub(r'!\[.*?\]\(.*?\)', '[Figure]', content)

            # Determine if this is abstract or keyword section
            is_abstract = False
            if content.startswith("摘要") or content.startswith("Abstract"):
                is_abstract = True
            if content.startswith("关键词") or content.startswith("Keywords"):
                is_abstract = True

            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Line spacing
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            para.paragraph_format.space_after = Pt(6)

            # First line indent for body text (not for references/abstract)
            if not in_references and not is_abstract:
                if not content.startswith("[") and not content.startswith("**"):
                    para.paragraph_format.first_line_indent = Cm(0.74)

            if in_references:
                # Hanging indent for references
                para.paragraph_format.left_indent = Cm(0.74)
                para.paragraph_format.first_line_indent = Cm(-0.74)
                add_runs_with_formatting(para, content, config, is_reference=True)
            elif is_abstract:
                add_runs_with_formatting(para, content, config, is_reference=True)
            else:
                add_runs_with_formatting(para, content, config)

    # Save document
    doc.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Citation integrity checker
# ---------------------------------------------------------------------------
def check_citation_integrity(md_text):
    """Check that all references in the list are cited in the text and vice versa."""
    ref_section_match = re.search(
        r'^#{1,3}\s*(?:参考文献|References)\s*$([\s\S]*)',
        md_text, re.MULTILINE | re.IGNORECASE
    )
    if not ref_section_match:
        return {
            "in_text_citations": [], "listed_refs": [],
            "cited_not_listed": [], "listed_not_cited": [], "is_consistent": False,
        }

    # Find all in-text citations [N] or [N-M] or [N,M,...]
    in_text_citations = set()
    body_text = md_text[:ref_section_match.start()]
    for match in re.finditer(r'\[([\d\-,\s]+)\]', body_text):
        group = match.group(1)
        # Parse range and list
        parts = group.split(",")
        for part in parts:
            part = part.strip()
            if "-" in part:
                nums = part.split("-")
                if len(nums) == 2 and nums[0].isdigit() and nums[1].isdigit():
                    start, end = int(nums[0]), int(nums[1])
                    for n in range(start, end + 1):
                        in_text_citations.add(n)
            elif part.isdigit():
                in_text_citations.add(int(part))

    # Find all references in the reference list (lines starting with [N]).
    listed_refs = set()
    ref_text = ref_section_match.group(1)
    for match in re.finditer(r'^\[(\d+)\]\s+', ref_text, re.MULTILINE):
        listed_refs.add(int(match.group(1)))

    cited_not_listed = in_text_citations - listed_refs
    listed_not_cited = listed_refs - in_text_citations

    return {
        "in_text_citations": sorted(in_text_citations),
        "listed_refs": sorted(listed_refs),
        "cited_not_listed": sorted(cited_not_listed),
        "listed_not_cited": sorted(listed_not_cited),
        "is_consistent": len(cited_not_listed) == 0 and len(listed_not_cited) == 0,
    }


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def process_file(input_path, output_path, lang, strict_citations=False, cn_font="楷体"):
    """Process a single markdown file into a docx file."""
    with open(input_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    # Override font for Chinese
    if lang == "zh" and cn_font in ("宋体", "楷体"):
        FONT_CONFIG["zh"]["body_font"] = cn_font

    # Check citation integrity
    integrity = check_citation_integrity(md_text)
    if not integrity["is_consistent"]:
        print(f"\n  WARNING: Citation integrity issues in {input_path}:")
        if integrity["cited_not_listed"]:
            print(f"    Cited in text but not in reference list: {integrity['cited_not_listed']}")
        if integrity["listed_not_cited"]:
            print(f"    In reference list but not cited in text: {integrity['listed_not_cited']}")
        print(f"    In-text citations: {len(integrity['in_text_citations'])}")
        print(f"    Listed references: {len(integrity['listed_refs'])}")
        if strict_citations:
            raise ValueError("Citation integrity validation failed")

    blocks = parse_markdown(md_text)
    generate_docx(blocks, output_path, lang)
    print(f"  Generated: {output_path}")
    print(f"  Blocks: {len(blocks)}")
    return output_path


FINAL_XLSX_HEADERS = ["Ref#", "Title", "Authors", "Journal", "Year", "Source DB",
                     "DOI", "PMID", "URL", "Cited By", "Open Access", "Abstract"]


def create_final_xlsx(md_path, selected_papers_path, output_path):
    """Create literature_search.xlsx with only actually cited references."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    md_text = Path(md_path).read_text(encoding="utf-8")
    integrity = check_citation_integrity(md_text)
    cited_nums = set(integrity["in_text_citations"])
    if not cited_nums:
        print("  WARNING: No citations found in text; cannot create final xlsx.")
        return None

    selected = json.loads(Path(selected_papers_path).read_text(encoding="utf-8"))
    ref_map = {r.get("ref_no"): r for r in selected if r.get("ref_no")}
    cited_records = []
    for ref_no in sorted(cited_nums):
        if ref_no in ref_map:
            cited_records.append(ref_map[ref_no])
        else:
            print(f"  WARNING: Cited ref [{ref_no}] not found in selected_papers.json")

    wb = Workbook()
    ws = wb.active
    ws.title = "Cited Literature"
    fill = PatternFill("solid", fgColor="2F5496")
    for column, header in enumerate(FINAL_XLSX_HEADERS, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for row_idx, record in enumerate(cited_records, 1):
        ref_no = record.get("ref_no", row_idx)
        values = [ref_no, record.get("title", ""), record.get("authors", ""),
                  record.get("journal", ""), record.get("year", ""),
                  record.get("source", ""), record.get("doi", ""),
                  record.get("pmid", ""), record.get("url", ""),
                  record.get("cited_by", 0), record.get("open_access", "Unknown"),
                  record.get("abstract", "")]
        for column, value in enumerate(values, 1):
            ws.cell(row_idx + 1, column, value)
    for col, width in {"A": 7, "B": 60, "C": 35, "D": 28, "E": 9, "F": 18,
                       "G": 30, "H": 12, "I": 45, "J": 12, "K": 14, "L": 70}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)
    print(f"  Final xlsx ({len(cited_records)} cited refs): {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(
        description="Generate formatted Word documents from markdown review files."
    )
    parser.add_argument("--input", help="Input markdown file path")
    parser.add_argument("--output", help="Output docx file path")
    parser.add_argument("--lang", choices=["zh", "en"], default="zh",
                        help="Language (zh=Chinese/楷体+Times New Roman, en=English/Times New Roman)")
    parser.add_argument("--cn-font", choices=["宋体", "楷体"], default="楷体",
                        help="Chinese body font (default: 楷体, alternative: 宋体)")
    parser.add_argument("--input-dir", help="Process both CN and EN files in directory")
    parser.add_argument("--strict-citations", action="store_true",
                        help="Refuse to create Word files when citation validation fails")
    parser.add_argument("--final-xlsx", action="store_true",
                        help="Create literature_search.xlsx with only cited refs (requires selected_papers.json in input-dir)")
    args = parser.parse_args()

    if args.input_dir:
        # Auto-detect and process both files
        input_dir = Path(args.input_dir)
        results = []

        cn_md = input_dir / "review_CN.md"
        cn_docx = input_dir / "review_CN.docx"
        en_md = input_dir / "review_EN.md"
        en_docx = input_dir / "review_EN.docx"

        if cn_md.exists():
            print(f"Processing Chinese: {cn_md}")
            process_file(str(cn_md), str(cn_docx), "zh", args.strict_citations, args.cn_font)
            results.append(str(cn_docx))
        else:
            print(f"WARNING: {cn_md} not found")

        if en_md.exists():
            print(f"Processing English: {en_md}")
            process_file(str(en_md), str(en_docx), "en", args.strict_citations)
            results.append(str(en_docx))
        else:
            print(f"WARNING: {en_md} not found")

        if args.final_xlsx:
            en_md_exists = en_md.exists()
            selected_path = input_dir / "selected_papers.json"
            if selected_path.exists() and en_md_exists:
                xlsx_path = input_dir / "literature_search.xlsx"
                create_final_xlsx(str(en_md), str(selected_path), str(xlsx_path))
            elif en_md_exists:
                print("  WARNING: --final-xlsx requires selected_papers.json in input-dir")
            else:
                print("  WARNING: --final-xlsx requires review_EN.md in input-dir")

        if results:
            print(f"\nDone! Generated {len(results)} file(s):")
            for r in results:
                print(f"  {r}")
        else:
            print("No markdown files found to process.")
            sys.exit(1)

    elif args.input and args.output:
        process_file(args.input, args.output, args.lang, args.strict_citations, args.cn_font if args.lang == "zh" else "楷体")

    else:
        parser.print_help()
        sys.exit(1)


if __name__ == "__main__":
    main()
