#!/usr/bin/env python3
"""Select the final cited literature set and export its audit files.

Use --keep-indices after screening candidate titles and abstracts. Without it,
the script ranks candidates by topic-keyword matches as a starting point only.

Reliability:
    Failures are reported with stable RG-xxx codes and mapped to exit codes
    (2 = input/usage, 1 = internal). --resume skips the run when the outputs are
    already up to date for the given --input, so an interrupted pipeline can be
    restarted cheaply; --force overrides that check.
"""

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

try:
    from _errors import (fail, run_guarded, summary_line, warn, write_report)
except ImportError:  # running from a copied location: fall back to sibling lookup
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from _errors import (fail, run_guarded, summary_line, warn, write_report)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    fail("RG-100", "openpyxl is required but not installed",
         "Install it with: pip install openpyxl")


STAMP_NAME = "selected_papers.stamp.json"


HEADERS = ["Ref#", "Title", "Authors", "Journal", "Year", "Source DB", "DOI",
           "PMID", "URL", "Cited By", "Open Access", "Abstract"]


def load_candidates(path: Path) -> list[dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail("RG-101", f"candidate file not found: {path}",
             "Run multi_db_search.py first, or point --input at the search JSON.")
    except OSError as exc:
        fail("RG-101", f"cannot read candidate file {path} ({exc})")
    except json.JSONDecodeError as exc:
        fail("RG-102", f"candidate file is not valid JSON: {path} ({exc})",
             "Re-run multi_db_search.py to regenerate the search JSON.")
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in ("unique_papers", "relevant_papers", "papers"):
            if isinstance(data.get(key), list):
                return data[key]
    fail("RG-102", f"expected a JSON list of papers in {path}",
         "Got a dict without unique_papers/relevant_papers/papers, or another type.")


def normalise(record: dict) -> dict:
    source = record.get("source") or record.get("database") or "Unknown"
    date = str(record.get("year") or record.get("publication_date") or "")
    doi = str(record.get("doi") or "").replace("https://doi.org/", "")
    authors = record.get("authors") or ""
    if isinstance(authors, list):
        authors = ", ".join(
            a.get("name", "") if isinstance(a, dict) else str(a) for a in authors
        )
    return {
        "title": str(record.get("title") or "").strip(),
        "authors": str(authors), "journal": str(record.get("journal") or ""),
        "year": date[:4], "source": str(source), "doi": doi,
        "pmid": str(record.get("pmid") or record.get("id") if source.lower() == "pubmed" else ""),
        "url": str(record.get("url") or (f"https://doi.org/{doi}" if doi else "")),
        "cited_by": int(record.get("cited_by") or record.get("citationCount") or 0),
        "open_access": "Yes" if record.get("has_fulltext") else "Unknown",
        "abstract": str(record.get("abstract") or "").strip(),
    }


def rank(records: list[dict], topic: str) -> list[dict]:
    terms = [t for t in re.findall(r"[\w-]+", topic.lower()) if len(t) > 2]
    def score(record: dict) -> tuple[int, int, str]:
        text = f"{record['title']} {record['abstract']}".lower()
        return (sum(t in text for t in terms), record["cited_by"], record["title"].lower())
    return sorted(records, key=score, reverse=True)


def write_excel(records: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Selected literature"
    fill = PatternFill("solid", fgColor="2F5496")
    for column, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for ref_no, record in enumerate(records, 1):
        values = [ref_no, record["title"], record["authors"], record["journal"], record["year"],
                  record["source"], record["doi"], record["pmid"], record["url"], record["cited_by"],
                  record["open_access"], record["abstract"]]
        for column, value in enumerate(values, 1):
            ws.cell(ref_no + 1, column, value)
    for col, width in {"A": 7, "B": 60, "C": 35, "D": 28, "E": 9, "F": 18,
                       "G": 30, "H": 12, "I": 45, "J": 12, "K": 14, "L": 70}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def _input_fingerprint(path: Path) -> str:
    if not path.is_file():
        fail("RG-101", f"candidate file not found: {path}",
             "Run multi_db_search.py first, or point --input at the search JSON.")
    try:
        digest = hashlib.sha256()
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 20), b""):
                digest.update(chunk)
        return digest.hexdigest()[:16]
    except OSError as exc:
        fail("RG-101", f"cannot read candidate file {path} ({exc})")


def _selection_signature(topic: str, keep_indices, max_papers: int) -> str:
    payload = json.dumps({"topic": topic, "keep": sorted(keep_indices or []),
                          "max": max_papers}, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def _outputs_are_current(output_dir: Path, fingerprint: str, signature: str) -> bool:
    stamp = output_dir / STAMP_NAME
    if not stamp.is_file():
        return False
    try:
        data = json.loads(stamp.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    if data.get("input_sha256") != fingerprint or data.get("signature") != signature:
        return False
    return all((output_dir / name).is_file()
               for name in ("selected_papers.json", "selected_papers.xlsx"))


def main_custom(input_path: str = None, output_dir: str = None,
                topic: str = "", keep_indices: list[int] = None,
                max_papers: int = 100, resume: bool = False,
                force: bool = False, error_report: str = "") -> None:
    if input_path and output_dir:
        _do_filter(input_path, output_dir, topic, keep_indices, max_papers,
                   resume=resume, force=force, error_report=error_report)
        return
    parser = argparse.ArgumentParser(description="Create the final cited literature set.")
    parser.add_argument("--input", required=True, help="Candidate literature_search.json")
    parser.add_argument("--output-dir", required=True, help="Topic output directory")
    parser.add_argument("--topic", default="", help="Used only for automatic ranking")
    parser.add_argument("--keep-indices", nargs="+", type=int,
                        help="1-based indices selected after manual screening")
    parser.add_argument("--max-papers", type=int, default=100)
    parser.add_argument("--resume", action="store_true",
                        help="Skip when selected_papers.json/.xlsx are already current for this --input")
    parser.add_argument("--force", action="store_true",
                        help="Re-run even when --resume would skip")
    parser.add_argument("--error-report", default="",
                        help="Write a JSON issue report to this path")
    args = parser.parse_args()
    _do_filter(args.input, args.output_dir, args.topic, args.keep_indices, args.max_papers,
               resume=args.resume, force=args.force, error_report=args.error_report)


def _do_filter(input_path, output_dir, topic, keep_indices, max_papers,
               resume=False, force=False, error_report=""):
    input_file = Path(input_path)
    output_path = Path(output_dir)
    fingerprint = _input_fingerprint(input_file)
    signature = _selection_signature(topic, keep_indices, max_papers)

    if resume and not force and _outputs_are_current(output_path, fingerprint, signature):
        print(f"Resume: selection already current for {input_file.name} "
              f"({STAMP_NAME} matches); nothing to do.")
        return

    # Validate max_papers range (used for auto-ranking when keep_indices is not provided)
    if keep_indices is None and not 1 <= max_papers <= 300:
        fail("RG-103", f"--max-papers must be between 1 and 300, got {max_papers}")

    candidates = [normalise(r) for r in load_candidates(input_file)]
    candidates = [r for r in candidates if r["title"]]
    if not candidates:
        fail("RG-102", f"no usable records with a title in {input_file}",
             "Check that the search step produced literature_search.json.")
    if keep_indices:
        bad = [i for i in keep_indices if i < 1 or i > len(candidates)]
        if bad:
            fail("RG-103", f"selected indices outside 1..{len(candidates)}: {bad}",
                 "Re-check the screening table before re-running.")
        selected = [candidates[i - 1] for i in keep_indices]
    else:
        selected = rank(candidates, topic)[:max_papers]
        warn("RG-202", "automatic ranking is only a starting point",
             "Screen titles and abstracts, then re-run with --keep-indices.")
    if len(selected) != len({(r["doi"].lower() or r["title"].lower()) for r in selected}):
        warn("RG-201", "auto-deduplicating duplicate titles or DOIs in the final selection")
        seen = set()
        deduped = []
        for r in selected:
            key = (r["doi"].lower() or r["title"].lower())
            if key not in seen:
                seen.add(key)
                deduped.append(r)
        selected = deduped

    output_path.mkdir(parents=True, exist_ok=True)
    for ref_no, record in enumerate(selected, 1):
        record["ref_no"] = ref_no
    (output_path / "selected_papers.json").write_text(
        json.dumps(selected, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_excel(selected, output_path / "selected_papers.xlsx")
    (output_path / STAMP_NAME).write_text(
        json.dumps({"input_sha256": fingerprint, "signature": signature,
                    "count": len(selected)}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Selected {len(selected)} papers")
    print(f"Saved: {output_path / 'selected_papers.json'}")
    print(f"Saved: {output_path / 'selected_papers.xlsx'}")

    if error_report:
        report = write_report(error_report)
        if report:
            print(f"Saved issue report: {report}")
    print(summary_line())


if __name__ == "__main__":
    run_guarded(main_custom)
